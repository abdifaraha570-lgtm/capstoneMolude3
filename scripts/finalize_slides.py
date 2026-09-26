"""Put the real run results into the Module 3 deck.

  * slide 8: replaces the two screenshot boxes with the GX images from reports/
  * slide 12: your GitHub URL
usage: python scripts/finalize_slides.py deck.pptx out.pptx https://github.com/you/repo
"""
import json
import sys
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
REP = ROOT / "reports"


def pick(*names):
    for n in names:
        if (REP / n).exists():
            return REP / n
    return None


def replace_in_runs(slide, old, new):
    done = False
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new); done = True
    return done


def put_image(slide, marker, img_path):
    box = next((s for s in slide.shapes if s.has_text_frame and marker in s.text_frame.text), None)
    if box is None or img_path is None:
        return False
    l, t, w, h = box.left, box.top, box.width, box.height
    for s in list(slide.shapes):          # remove the dashed frame + placeholder text
        if s.left == l and s.top == t and s.width == w and s.height == h:
            s._element.getparent().remove(s._element)
    iw, ih = Image.open(img_path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(img_path), l + (w - pw) // 2, t + (h - ph) // 2, pw, ph)
    return True


def main(deck, out, url):
    prs = Presentation(deck)
    s8, s12 = prs.slides[7], prs.slides[11]
    raw_img = pick("gx_raw_screenshot.png", "gx_raw_results.png")
    proc_img = pick("gx_processed_screenshot.png", "gx_processed_results.png")
    print("slide 8 raw:", put_image(s8, "raw suite results", raw_img), raw_img)
    print("slide 8 processed:", put_image(s8, "processed suite results", proc_img), proc_img)

    print("slide 12:", replace_in_runs(s12, "https://github.com/<your-username>/jjushycsh-noshow-analytics", url))
    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    main(*sys.argv[1:4])
